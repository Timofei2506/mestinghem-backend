"""
Document renderers — deterministic, template driven.

A validated *spec* (plain JSON) goes in, a real Office file comes out.
Rendering never depends on the language model, so output is reproducible and
always well-formed even if inference is unavailable.

Templates follow the house style: obsidian background, champagne-gold accents,
generous whitespace, one idea per slide.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import datetime, timezone
from typing import Any, Iterable, Sequence

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor as DocxRGB
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt as PPt

log = logging.getLogger("documents")

# ---------------------------------------------------------------------------
# House palette
# ---------------------------------------------------------------------------
OBSIDIAN = RGBColor(0x0A, 0x0B, 0x10)
GRAPHITE = RGBColor(0x14, 0x16, 0x1F)
SLATE = RGBColor(0x1E, 0x22, 0x2E)
GOLD = RGBColor(0xC8, 0xA4, 0x5C)
GOLD_LIGHT = RGBColor(0xE8, 0xC8, 0x7E)
IVORY = RGBColor(0xF3, 0xF1, 0xEC)
MUTED = RGBColor(0x9A, 0xA0, 0xB0)
ACCENT_TEAL = RGBColor(0x4F, 0xC3, 0xA1)
ACCENT_ROSE = RGBColor(0xD9, 0x6A, 0x7A)

HEX = {
    "obsidian": "0A0B10",
    "graphite": "14161F",
    "slate": "1E222E",
    "gold": "C8A45C",
    "gold_light": "E8C87E",
    "ivory": "F3F1EC",
    "muted": "9AA0B0",
}

BODY_FONT = "Segoe UI"
HEAD_FONT = "Segoe UI Semibold"
MONO_FONT = "Consolas"

SLIDE_W = Emu(12_192_000)   # 16:9
SLIDE_H = Emu(6_858_000)


def _today() -> str:
    return datetime.now(timezone.utc).strftime("%d.%m.%Y")


def _truncate(text: Any, limit: int = 400) -> str:
    value = str(text or "").strip()
    return value if len(value) <= limit else value[: limit - 1].rstrip() + "…"


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------
def _fill(shape: Any, color: RGBColor, *, line: RGBColor | None = None) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = PPt(1)
    shape.shadow.inherit = False


def _text(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: int = 18,
    color: RGBColor = IVORY,
    bold: bool = False,
    align: Any = PP_ALIGN.LEFT,
    font: str = BODY_FONT,
    anchor: Any = MSO_ANCHOR.TOP,
    spacing: float = 1.15,
) -> Any:
    box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = spacing
        run = paragraph.add_run()
        run.text = line
        run.font.size = PPt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font
    return box


def _rect(slide: Any, left: float, top: float, width: float, height: float, color: RGBColor, shape_type: Any = MSO_SHAPE.RECTANGLE) -> Any:
    shape = slide.shapes.add_shape(shape_type, Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    _fill(shape, color)
    return shape


def _blank(prs: Presentation) -> Any:
    return prs.slides.add_slide(prs.slide_layouts[6])


def _chrome(slide: Any, *, page: int | None = None, total: int | None = None, brand: str = "Mestinghem AI") -> None:
    """Background, hairline accents and footer — shared by every slide."""
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, OBSIDIAN)
    _rect(slide, 0, 0, SLIDE_W, Emu(28_000), GOLD)
    _rect(slide, Emu(700_000), Emu(6_420_000), Emu(1_100_000), Emu(9_000), GOLD)
    _text(
        slide, Emu(700_000), Emu(6_470_000), Emu(6_000_000), Emu(280_000),
        brand.upper(), size=9, color=MUTED, bold=True,
    )
    if page is not None:
        label = f"{page:02d}" + (f" / {total:02d}" if total else "")
        _text(
            slide, Emu(10_600_000), Emu(6_470_000), Emu(900_000), Emu(280_000),
            label, size=9, color=MUTED, align=PP_ALIGN.RIGHT,
        )


def _kicker(slide: Any, text: str, top: float = 620_000) -> None:
    if not text:
        return
    _rect(slide, Emu(700_000), Emu(top), Emu(300_000), Emu(28_000), GOLD)
    _text(
        slide, Emu(700_000), Emu(top + 90_000), Emu(8_000_000), Emu(300_000),
        str(text).upper(), size=11, color=GOLD_LIGHT, bold=True,
    )


def _bullets(slide: Any, items: Sequence[str], *, left: float = 700_000, top: float = 1_900_000,
             width: float = 10_800_000, size: int = 16, gap: int = 420_000) -> None:
    cursor = top
    for item in items:
        text = _truncate(item, 300)
        if not text:
            continue
        _rect(slide, Emu(left), Emu(cursor + 120_000), Emu(70_000), Emu(70_000), GOLD, MSO_SHAPE.OVAL)
        box = _text(
            slide, Emu(left + 260_000), Emu(cursor), Emu(width - 260_000), Emu(gap),
            text, size=size, color=IVORY, spacing=1.2,
        )
        lines = max(1, (len(text) // 78) + 1)
        cursor += int(gap * 0.72 * lines)


def _stat_cards(slide: Any, stats: Sequence[dict[str, Any]], *, top: float = 1_900_000) -> None:
    cards = list(stats)[:4]
    if not cards:
        return
    margin = 700_000
    gap = 260_000
    width = int((SLIDE_W - 2 * margin - gap * (len(cards) - 1)) / len(cards))
    for index, card in enumerate(cards):
        left = margin + index * (width + gap)
        panel = _rect(slide, Emu(left), Emu(top), Emu(width), Emu(1_750_000), GRAPHITE)
        _rect(slide, Emu(left), Emu(top), Emu(width), Emu(30_000), GOLD)
        _text(
            slide, Emu(left + 220_000), Emu(top + 260_000), Emu(width - 440_000), Emu(620_000),
            _truncate(card.get("value"), 14), size=34, color=GOLD_LIGHT, bold=True, font=HEAD_FONT,
        )
        _text(
            slide, Emu(left + 220_000), Emu(top + 980_000), Emu(width - 440_000), Emu(620_000),
            _truncate(card.get("label"), 70), size=12, color=MUTED,
        )
        del panel


def _chart_slide(prs: Presentation, spec: dict[str, Any], *, page: int, total: int, brand: str) -> None:
    slide = _blank(prs)
    _chrome(slide, page=page, total=total, brand=brand)
    _kicker(slide, spec.get("kicker") or "Данные")
    _text(
        slide, Emu(700_000), Emu(900_000), Emu(10_800_000), Emu(820_000),
        _truncate(spec.get("title") or "Показатели", 120), size=30, color=IVORY, bold=True, font=HEAD_FONT,
    )
    chart_spec = spec.get("chart") or {}
    categories = [str(c) for c in (chart_spec.get("categories") or [])][:12]
    series = chart_spec.get("series") or []
    if categories and series:
        data = CategoryChartData()
        data.categories = categories
        for item in series[:4]:
            values = [float(v or 0) for v in (item.get("values") or [])][: len(categories)]
            values += [0.0] * (len(categories) - len(values))
            data.add_series(_truncate(item.get("name"), 40) or "Ряд", values)
        kind = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "area": XL_CHART_TYPE.AREA,
        }.get(str(chart_spec.get("type", "bar")).lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)
        graphic = slide.shapes.add_chart(
            kind, Emu(700_000), Emu(1_950_000), Emu(10_800_000), Emu(4_100_000), data
        )
        chart = graphic.chart
        chart.has_legend = len(series) > 1 or kind == XL_CHART_TYPE.PIE
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = PPt(11)
            chart.legend.font.color.rgb = MUTED
        plot = chart.plots[0]
        plot.gap_width = 90
        palette = [GOLD, ACCENT_TEAL, ACCENT_ROSE, GOLD_LIGHT]
        for index, serie in enumerate(plot.series):
            serie.format.fill.solid()
            serie.format.fill.fore_color.rgb = palette[index % len(palette)]
        axis = chart.category_axis
        axis.tick_labels.font.size = PPt(10)
        axis.tick_labels.font.color.rgb = MUTED
        axis.format.line.color.rgb = SLATE
        value_axis = chart.value_axis
        value_axis.tick_labels.font.size = PPt(10)
        value_axis.tick_labels.font.color.rgb = MUTED
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = SLATE
        value_axis.format.line.color.rgb = SLATE
    if spec.get("note"):
        _text(
            slide, Emu(700_000), Emu(6_050_000), Emu(10_800_000), Emu(300_000),
            _truncate(spec.get("note"), 160), size=11, color=MUTED,
        )


def render_pptx(spec: dict[str, Any], *, brand: str = "Mestinghem AI") -> bytes:
    """Build a 16:9 deck from a validated spec."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slides: list[dict[str, Any]] = [s for s in (spec.get("slides") or []) if isinstance(s, dict)]
    total = max(len(slides), 1)
    title = _truncate(spec.get("title") or "Презентация", 140)
    subtitle = _truncate(spec.get("subtitle") or "", 220)
    author = _truncate(spec.get("author") or brand, 80)

    # --- cover ------------------------------------------------------------
    slide = _blank(prs)
    _chrome(slide, brand=brand)
    _rect(slide, Emu(9_200_000), Emu(-900_000), Emu(4_200_000), Emu(4_200_000), GRAPHITE, MSO_SHAPE.OVAL)
    _rect(slide, Emu(10_400_000), Emu(3_600_000), Emu(2_600_000), Emu(2_600_000), SLATE, MSO_SHAPE.OVAL)
    _rect(slide, Emu(700_000), Emu(2_050_000), Emu(900_000), Emu(40_000), GOLD)
    _text(
        slide, Emu(700_000), Emu(2_320_000), Emu(9_600_000), Emu(1_800_000),
        title, size=44, color=IVORY, bold=True, font=HEAD_FONT, spacing=1.05,
    )
    if subtitle:
        _text(slide, Emu(700_000), Emu(4_250_000), Emu(9_000_000), Emu(900_000), subtitle, size=17, color=MUTED)
    _text(
        slide, Emu(700_000), Emu(5_500_000), Emu(9_000_000), Emu(400_000),
        f"{author}   ·   {_today()}", size=12, color=GOLD_LIGHT, bold=True,
    )

    # --- agenda -----------------------------------------------------------
    agenda = [
        _truncate(s.get("title") or s.get("heading") or "", 90)
        for s in slides
        if s.get("layout") != "section"
    ]
    agenda = [a for a in agenda if a][:9]
    if agenda:
        slide = _blank(prs)
        _chrome(slide, page=1, total=total, brand=brand)
        _kicker(slide, spec.get("agenda_kicker") or "Содержание")
        _text(
            slide, Emu(700_000), Emu(900_000), Emu(10_800_000), Emu(760_000),
            "О чём пойдёт речь", size=30, color=IVORY, bold=True, font=HEAD_FONT,
        )
        cursor = 1_950_000
        for index, item in enumerate(agenda, start=1):
            _text(
                slide, Emu(700_000), Emu(cursor), Emu(700_000), Emu(380_000),
                f"{index:02d}", size=15, color=GOLD, bold=True, font=HEAD_FONT,
            )
            _text(slide, Emu(1_500_000), Emu(cursor), Emu(9_800_000), Emu(380_000), item, size=16, color=IVORY)
            _rect(slide, Emu(1_500_000), Emu(cursor + 330_000), Emu(9_400_000), Emu(6_000), SLATE)
            cursor += 470_000

    # --- body -------------------------------------------------------------
    for index, item in enumerate(slides, start=1):
        layout = str(item.get("layout") or "content").lower()
        heading = _truncate(item.get("title") or item.get("heading") or "", 120)
        points = [_truncate(p, 300) for p in (item.get("points") or []) if str(p).strip()][:6]

        if layout == "section":
            slide = _blank(prs)
            _chrome(slide, page=index, total=total, brand=brand)
            _rect(slide, 0, Emu(2_600_000), SLIDE_W, Emu(1_500_000), GRAPHITE)
            _rect(slide, Emu(700_000), Emu(3_150_000), Emu(700_000), Emu(34_000), GOLD)
            _text(
                slide, Emu(700_000), Emu(3_330_000), Emu(10_000_000), Emu(700_000),
                heading, size=34, color=IVORY, bold=True, font=HEAD_FONT,
            )
            if item.get("note"):
                _text(
                    slide, Emu(700_000), Emu(4_250_000), Emu(9_500_000), Emu(500_000),
                    _truncate(item.get("note"), 200), size=14, color=MUTED,
                )
            continue

        slide = _blank(prs)
        _chrome(slide, page=index, total=total, brand=brand)
        _kicker(slide, _truncate(item.get("kicker") or "", 60))
        _text(
            slide, Emu(700_000), Emu(900_000), Emu(10_800_000), Emu(820_000),
            heading or "—", size=30, color=IVORY, bold=True, font=HEAD_FONT,
        )

        stats = item.get("stats") or []
        if layout in ("stats", "kpi") and stats:
            _stat_cards(slide, stats, top=1_900_000)
            if points:
                _bullets(slide, points, top=3_950_000, size=14, gap=380_000)
        elif layout == "two_column":
            columns = item.get("columns") or [{"title": "", "points": points}]
            width = 5_100_000
            for col_index, column in enumerate(columns[:2]):
                left = 700_000 + col_index * (width + 500_000)
                _rect(slide, Emu(left), Emu(1_900_000), Emu(width), Emu(4_000_000), GRAPHITE)
                _rect(slide, Emu(left), Emu(1_900_000), Emu(width), Emu(28_000), GOLD)
                _text(
                    slide, Emu(left + 260_000), Emu(2_150_000), Emu(width - 520_000), Emu(420_000),
                    _truncate(column.get("title"), 70), size=17, color=GOLD_LIGHT, bold=True, font=HEAD_FONT,
                )
                _bullets(
                    slide,
                    [_truncate(p, 220) for p in (column.get("points") or []) if str(p).strip()][:6],
                    left=left + 260_000, top=2_750_000, width=width - 520_000, size=14, gap=380_000,
                )
        elif layout == "quote":
            _rect(slide, Emu(700_000), Emu(2_000_000), Emu(60_000), Emu(2_400_000), GOLD)
            _text(
                slide, Emu(1_100_000), Emu(2_050_000), Emu(10_000_000), Emu(2_300_000),
                f"«{_truncate(item.get('quote') or heading, 300)}»", size=26, color=IVORY, font=HEAD_FONT, spacing=1.25,
            )
            if item.get("author"):
                _text(
                    slide, Emu(1_100_000), Emu(4_500_000), Emu(9_000_000), Emu(400_000),
                    f"— {_truncate(item.get('author'), 80)}", size=14, color=GOLD_LIGHT, bold=True,
                )
        elif layout == "table":
            rows = item.get("rows") or []
            headers = item.get("headers") or []
            if headers and rows:
                cols = min(len(headers), 5)
                rows = rows[:7]
                width = 10_800_000
                left, top = 700_000, 1_950_000
                row_h = 460_000
                col_w = width // cols
                for r, row in enumerate([headers] + rows):
                    for c in range(cols):
                        cell_left = left + c * col_w
                        cell_top = top + r * row_h
                        bg = SLATE if r == 0 else (GRAPHITE if r % 2 else OBSIDIAN)
                        _rect(slide, Emu(cell_left), Emu(cell_top), Emu(col_w - 20_000), Emu(row_h - 20_000), bg)
                        _text(
                            slide, Emu(cell_left + 140_000), Emu(cell_top + 110_000), Emu(col_w - 300_000), Emu(row_h - 100_000),
                            _truncate(row[c] if c < len(row) else "", 60),
                            size=12 if r else 12,
                            color=GOLD_LIGHT if r == 0 else IVORY,
                            bold=(r == 0),
                        )
        else:
            if item.get("lead"):
                _text(
                    slide, Emu(700_000), Emu(1_750_000), Emu(10_800_000), Emu(560_000),
                    _truncate(item.get("lead"), 240), size=15, color=GOLD_LIGHT,
                )
            _bullets(slide, points, top=2_450_000 if item.get("lead") else 1_950_000)
            if stats:
                _stat_cards(slide, stats, top=5_050_000)
        if item.get("note"):
            _text(
                slide, Emu(700_000), Emu(6_050_000), Emu(10_800_000), Emu(300_000),
                _truncate(item.get("note"), 170), size=11, color=MUTED,
            )
        if item.get("chart"):
            _chart_slide(prs, {"title": item.get("chart_title") or heading, "kicker": "Данные",
                               "chart": item["chart"], "note": item.get("chart_note")},
                         page=index, total=total, brand=brand)

    # --- closing ----------------------------------------------------------
    slide = _blank(prs)
    _chrome(slide, page=total, total=total, brand=brand)
    _rect(slide, Emu(700_000), Emu(2_500_000), Emu(700_000), Emu(34_000), GOLD)
    _text(
        slide, Emu(700_000), Emu(2_700_000), Emu(10_000_000), Emu(900_000),
        _truncate(spec.get("closing_title") or "Спасибо за внимание", 90),
        size=36, color=IVORY, bold=True, font=HEAD_FONT,
    )
    _text(
        slide, Emu(700_000), Emu(3_750_000), Emu(10_000_000), Emu(700_000),
        _truncate(spec.get("closing_text") or "Готов ответить на вопросы и перейти к следующим шагам.", 200),
        size=16, color=MUTED,
    )
    contacts = spec.get("contacts") or []
    if contacts:
        _text(
            slide, Emu(700_000), Emu(4_700_000), Emu(10_000_000), Emu(400_000),
            "   ·   ".join(_truncate(c, 60) for c in contacts[:4]), size=13, color=GOLD_LIGHT, bold=True,
        )

    core = prs.core_properties
    core.title = title
    core.author = author
    core.subject = _truncate(spec.get("subtitle") or "", 120)
    core.comments = f"Generated by {brand}"
    core.category = "presentation"

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
_XL_FONT = "Calibri"


def _xl_header(cell: Any) -> None:
    cell.font = Font(name=_XL_FONT, size=11, bold=True, color="FF" + HEX["gold_light"])
    cell.fill = PatternFill("solid", fgColor=HEX["slate"])
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="thin", color=HEX["gold"]))


def _xl_body(cell: Any, *, index: int = 0, align: str = "left") -> None:
    cell.font = Font(name=_XL_FONT, size=11, color="FF" + HEX["ivory"])
    cell.fill = PatternFill("solid", fgColor=HEX["graphite"] if index % 2 == 0 else HEX["obsidian"])
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="hair", color="2A2F3D"))


def _xl_cover(wb: Workbook, spec: dict[str, Any], *, brand: str) -> None:
    ws = wb.create_sheet("Обзор", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = HEX["gold"]
    for column, width in zip("ABCDEFGH", (3, 34, 26, 20, 20, 20, 20, 3)):
        ws.column_dimensions[column].width = width
    ws["B2"] = brand.upper()
    ws["B2"].font = Font(name=_XL_FONT, size=11, bold=True, color=HEX["gold"])
    ws["B4"] = _truncate(spec.get("title") or "Рабочая книга", 120)
    ws["B4"].font = Font(name=_XL_FONT, size=26, bold=True, color="FF" + HEX["ivory"])
    ws["B5"] = _truncate(spec.get("subtitle") or "", 200)
    ws["B5"].font = Font(name=_XL_FONT, size=12, color="FF" + HEX["muted"])
    ws["B7"] = f"Дата: {_today()}"
    ws["B7"].font = Font(name=_XL_FONT, size=10, color="FF" + HEX["muted"])
    ws["B8"] = f"Листов: {len(wb.sheetnames)}"
    ws["B8"].font = Font(name=_XL_FONT, size=10, color="FF" + HEX["muted"])

    kpis = spec.get("kpis") or []
    row = 10
    if kpis:
        ws.cell(row=row, column=2, value="Ключевые показатели").font = Font(
            name=_XL_FONT, size=13, bold=True, color="FF" + HEX["gold_light"]
        )
        row += 1
        for index, kpi in enumerate(kpis[:6]):
            label = ws.cell(row=row + index, column=2, value=_truncate(kpi.get("label"), 80))
            value = ws.cell(row=row + index, column=3, value=kpi.get("value"))
            label.font = Font(name=_XL_FONT, size=11, color="FF" + HEX["ivory"])
            value.font = Font(name=_XL_FONT, size=12, bold=True, color="FF" + HEX["gold_light"])
            value.alignment = Alignment(horizontal="right")
            if isinstance(kpi.get("value"), (int, float)):
                value.number_format = kpi.get("format") or "#,##0.00"
        row += len(kpis[:6]) + 2

    notes = spec.get("notes") or []
    if notes:
        ws.cell(row=row, column=2, value="Комментарии").font = Font(
            name=_XL_FONT, size=13, bold=True, color="FF" + HEX["gold_light"]
        )
        for index, note in enumerate(notes[:12]):
            cell = ws.cell(row=row + 1 + index, column=2, value=_truncate(note, 300))
            cell.font = Font(name=_XL_FONT, size=11, color="FF" + HEX["muted"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row + 1 + index, start_column=2, end_row=row + 1 + index, end_column=6)
            ws.row_dimensions[row + 1 + index].height = 30
    for row_index in range(1, row + 20):
        ws.row_dimensions[row_index].height = ws.row_dimensions[row_index].height or 20
    ws.row_dimensions[4].height = 34
    ws.sheet_view.zoomScale = 110


def _xl_data_sheet(wb: Workbook, sheet: dict[str, Any], *, name: str) -> None:
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = HEX["slate"]
    headers = [str(h) for h in (sheet.get("headers") or [])][:20]
    rows = sheet.get("rows") or []
    if not headers:
        ws["B2"] = _truncate(sheet.get("title") or name, 120)
        ws["B2"].font = Font(name=_XL_FONT, size=16, bold=True, color="FF" + HEX["ivory"])
        return

    start_row = 3
    ws.cell(row=1, column=1, value=_truncate(sheet.get("title") or name, 120)).font = Font(
        name=_XL_FONT, size=15, bold=True, color="FF" + HEX["ivory"]
    )
    if sheet.get("subtitle"):
        ws.cell(row=2, column=1, value=_truncate(sheet.get("subtitle"), 180)).font = Font(
            name=_XL_FONT, size=10, color="FF" + HEX["muted"]
        )

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        _xl_header(cell)
        widths = sheet.get("widths") or []
        width = widths[col_index - 1] if col_index - 1 < len(widths) else None
        ws.column_dimensions[get_column_letter(col_index)].width = float(width or max(14, min(42, len(header) + 8)))
    ws.row_dimensions[start_row].height = 28

    numeric_columns: set[int] = set()
    for row_index, row in enumerate(rows[:5000], start=start_row + 1):
        values = list(row) if isinstance(row, (list, tuple)) else [row]
        for col_index in range(1, len(headers) + 1):
            raw = values[col_index - 1] if col_index - 1 < len(values) else None
            if isinstance(raw, str):
                cleaned = raw.strip().replace(" ", "").replace("\u00a0", "")
                try:
                    raw = float(cleaned.replace(",", ".")) if re.match(r"^-?\d+([.,]\d+)?%?$", cleaned) and cleaned[-1] != "%" else raw
                except (ValueError, AttributeError):
                    pass
            cell = ws.cell(row=row_index, column=col_index, value=raw)
            _xl_body(cell, index=row_index - start_row)
            if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                numeric_columns.add(col_index)
                cell.number_format = sheet.get("number_format") or "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.number_format = "@"
        ws.row_dimensions[row_index].height = 20

    last_row = start_row + len(rows[:5000])
    if rows:
        ref = f"A{start_row}:{get_column_letter(len(headers))}{last_row}"
        table = Table(displayName=f"tbl_{''.join(ch for ch in name if ch.isalnum())[:20] or 'data'}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        try:
            ws.add_table(table)
        except Exception:
            pass
        ws.auto_filter.ref = ref
        # Totals row
        total_row = last_row + 1
        ws.cell(row=total_row, column=1, value="Итого").font = Font(
            name=_XL_FONT, size=11, bold=True, color="FF" + HEX["gold_light"]
        )
        for col_index in sorted(numeric_columns):
            letter = get_column_letter(col_index)
            cell = ws.cell(
                row=total_row, column=col_index,
                value=f"=SUBTOTAL(9,{letter}{start_row + 1}:{letter}{last_row})",
            )
            cell.font = Font(name=_XL_FONT, size=11, bold=True, color="FF" + HEX["gold_light"])
            cell.fill = PatternFill("solid", fgColor=HEX["slate"])
            cell.number_format = sheet.get("number_format") or "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
            if sheet.get("data_bars", True):
                ws.conditional_formatting.add(
                    f"{letter}{start_row + 1}:{letter}{last_row}",
                    DataBarRule(start_type="num", start_value=0, end_type="max", color=HEX["gold"], showValue=True),
                )
        if sheet.get("color_scale") and len(numeric_columns) >= 1:
            letter = get_column_letter(sorted(numeric_columns)[0])
            ws.conditional_formatting.add(
                f"{letter}{start_row + 1}:{letter}{last_row}",
                ColorScaleRule(start_type="min", start_color="3A2B12", mid_type="percentile",
                               mid_value=50, mid_color=HEX["gold"], end_type="max", end_color=HEX["gold_light"]),
            )
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.sheet_view.zoomScale = 100

    for chart in (sheet.get("charts") or [])[:3]:
        _xl_chart(ws, chart, headers=headers, start_row=start_row, rows=len(rows), numeric_columns=numeric_columns)


def _xl_chart(ws: Any, chart: dict[str, Any], *, headers: Sequence[str], start_row: int,
              rows: int, numeric_columns: Iterable[int]) -> None:
    if rows < 2:
        return
    kind = str(chart.get("type", "bar")).lower()
    instance = {"line": LineChart(), "pie": PieChart()}.get(kind, BarChart())
    instance.style = 2
    instance.title = _truncate(chart.get("title") or "Динамика", 80)
    instance.height = chart.get("height") or 8.5
    instance.width = chart.get("width") or 17
    numeric = sorted(numeric_columns)
    columns = chart.get("columns") or numeric[:2]
    label_col = chart.get("label_column") or 1
    for col_index in columns:
        try:
            col_index = int(col_index)
        except (TypeError, ValueError):
            continue
        if not 1 <= col_index <= len(headers):
            continue
        data = Reference(ws, min_col=col_index, min_row=start_row, max_row=start_row + rows)
        instance.add_data(data, titles_from_data=True)
    labels = Reference(ws, min_col=label_col, min_row=start_row + 1, max_row=start_row + rows)
    instance.set_categories(labels)
    anchor = chart.get("anchor") or f"A{start_row + rows + 4}"
    try:
        ws.add_chart(instance, anchor)
    except Exception as exc:  # pragma: no cover
        log.debug("chart skipped: %s", exc)


def render_xlsx(spec: dict[str, Any], *, brand: str = "Mestinghem AI") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    sheets = spec.get("sheets") or []
    for index, sheet in enumerate(sheets[:12]):
        if not isinstance(sheet, dict):
            continue
        name = _truncate(sheet.get("name") or f"Лист {index + 1}", 28)
        _xl_data_sheet(wb, sheet, name=name)
    if not wb.sheetnames:
        _xl_data_sheet(wb, {"title": spec.get("title"), "headers": ["Показатель", "Значение"], "rows": []}, name="Данные")
    _xl_cover(wb, spec, brand=brand)
    wb.properties.title = _truncate(spec.get("title") or "Workbook", 120)
    wb.properties.creator = brand
    wb.properties.company = brand
    wb.properties.description = _truncate(spec.get("subtitle") or "", 200)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _docx_shade(paragraph: Any, hexcolor: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:fill"), hexcolor)
    properties.append(shade)


def _docx_border(paragraph: Any, *, color: str = HEX["gold"], size: int = 12, position: str = "bottom") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    edge = OxmlElement(f"w:{position}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "6")
    edge.set(qn("w:color"), color)
    borders.append(edge)


def _docx_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


def render_docx(spec: dict[str, Any], *, brand: str = "Mestinghem AI") -> bytes:
    doc = DocxDocument()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DocxRGB(0x1B, 0x1E, 0x27)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, space_before in (
        ("Heading 1", 20, DocxRGB(0x1A, 0x1A, 0x1A), 18),
        ("Heading 2", 15, DocxRGB(0x8A, 0x6D, 0x2F), 14),
        ("Heading 3", 12.5, DocxRGB(0x33, 0x38, 0x45), 10),
    ):
        style = styles[name]
        style.font.name = "Calibri Light"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(6)

    # --- cover ------------------------------------------------------------
    kicker = doc.add_paragraph()
    run = kicker.add_run(brand.upper())
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = DocxRGB(0x8A, 0x6D, 0x2F)
    _docx_border(kicker, color=HEX["gold"], size=6)

    title = doc.add_paragraph()
    title_run = title.add_run(_truncate(spec.get("title") or "Документ", 160))
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.name = "Calibri Light"
    title.paragraph_format.space_after = Pt(4)

    if spec.get("subtitle"):
        subtitle = doc.add_paragraph()
        sub_run = subtitle.add_run(_truncate(spec.get("subtitle"), 220))
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = DocxRGB(0x5A, 0x60, 0x6E)
        subtitle.paragraph_format.space_after = Pt(16)

    meta_bits = [b for b in (spec.get("author") or brand, _today(), spec.get("doc_type")) if b]
    if meta_bits:
        meta = doc.add_paragraph()
        meta_run = meta.add_run("   ·   ".join(str(b) for b in meta_bits))
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = DocxRGB(0x8A, 0x90, 0x9E)
        _docx_border(meta, color="D8D2C4", size=4)
        meta.paragraph_format.space_after = Pt(18)

    if spec.get("summary"):
        lead = doc.add_paragraph()
        lead_run = lead.add_run(_truncate(spec.get("summary"), 900))
        lead_run.font.size = Pt(11.5)
        lead_run.italic = True
        lead_run.font.color.rgb = DocxRGB(0x33, 0x38, 0x45)
        _docx_shade(lead, "F7F4EC")
        lead.paragraph_format.space_before = Pt(6)
        lead.paragraph_format.space_after = Pt(14)

    # --- table of contents ------------------------------------------------
    if spec.get("toc", True) and (spec.get("sections") or []):
        heading = doc.add_paragraph()
        heading_run = heading.add_run("Содержание")
        heading_run.font.size = Pt(12)
        heading_run.font.bold = True
        heading_run.font.color.rgb = DocxRGB(0x8A, 0x6D, 0x2F)
        toc = doc.add_paragraph()
        _docx_field(toc, r'TOC \o "1-2" \h \z \u')
        note = doc.add_paragraph()
        note_run = note.add_run("Обновите поле в Word: ПКМ → «Обновить поле».")
        note_run.font.size = Pt(8.5)
        note_run.font.color.rgb = DocxRGB(0x9A, 0xA0, 0xB0)
        note.paragraph_format.space_after = Pt(14)
        doc.add_page_break()

    # --- body -------------------------------------------------------------
    for item in spec.get("sections") or []:
        if not isinstance(item, dict):
            continue
        level = int(item.get("level") or 1)
        if item.get("page_break"):
            doc.add_page_break()
        if item.get("heading"):
            doc.add_heading(_truncate(item.get("heading"), 160), level=min(level, 3))
        for block in item.get("blocks") or []:
            if isinstance(block, str):
                block = {"type": "paragraph", "text": block}
            kind = str(block.get("type") or "paragraph").lower()
            if kind == "paragraph" and block.get("text"):
                doc.add_paragraph(_truncate(block.get("text"), 3000))
            elif kind in ("bullet", "bullets"):
                entries = block.get("items") or ([block.get("text")] if block.get("text") else [])
                for entry in entries:
                    if str(entry).strip():
                        doc.add_paragraph(_truncate(entry, 600), style="List Bullet")
            elif kind in ("numbered", "steps"):
                entries = block.get("items") or ([block.get("text")] if block.get("text") else [])
                for entry in entries:
                    if str(entry).strip():
                        doc.add_paragraph(_truncate(entry, 600), style="List Number")
            elif kind == "quote" and block.get("text"):
                quote = doc.add_paragraph()
                quote_run = quote.add_run(_truncate(block.get("text"), 800))
                quote_run.italic = True
                quote_run.font.color.rgb = DocxRGB(0x33, 0x38, 0x45)
                _docx_border(quote, color=HEX["gold"], size=18, position="left")
                quote.paragraph_format.left_indent = Cm(0.6)
            elif kind == "table":
                headers = block.get("headers") or []
                rows = block.get("rows") or []
                if not headers:
                    continue
                table = doc.add_table(rows=1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Light Grid Accent 1"
                for index, header in enumerate(headers):
                    cell = table.rows[0].cells[index]
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    run = paragraph.add_run(_truncate(header, 60))
                    run.font.bold = True
                    run.font.size = Pt(10)
                    _docx_shade(paragraph, HEX["slate"])
                for row in rows[:200]:
                    cells = table.add_row().cells
                    for index, value in enumerate(list(row)[: len(headers)]):
                        cells[index].text = ""
                        paragraph = cells[index].paragraphs[0]
                        run = paragraph.add_run(_truncate(value, 200))
                        run.font.size = Pt(10)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            elif kind == "callout" and block.get("text"):
                callout = doc.add_paragraph()
                run = callout.add_run(_truncate(block.get("text"), 600))
                run.font.size = Pt(10.5)
                run.font.bold = bool(block.get("bold", True))
                _docx_shade(callout, "FBF6EA")
                _docx_border(callout, color=HEX["gold"], size=8, position="left")
                callout.paragraph_format.left_indent = Cm(0.4)

    # --- footer with page numbers -----------------------------------------
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"{brand}  ·  ")
    run.font.size = Pt(8)
    run.font.color.rgb = DocxRGB(0x9A, 0xA0, 0xB0)
    _docx_field(paragraph, "PAGE")
    header = section.header
    header_run = header.paragraphs[0].add_run(_truncate(spec.get("title") or brand, 90))
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = DocxRGB(0x9A, 0xA0, 0xB0)

    core = doc.core_properties
    core.title = _truncate(spec.get("title") or "Document", 120)
    core.author = _truncate(spec.get("author") or brand, 80)
    core.comments = f"Generated by {brand}"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Markdown / HTML
# ---------------------------------------------------------------------------
def render_markdown(spec: dict[str, Any]) -> str:
    lines = [f"# {spec.get('title') or 'Документ'}", ""]
    if spec.get("subtitle"):
        lines += [f"> {spec['subtitle']}", ""]
    lines += [f"*{spec.get('author') or 'Mestinghem AI'} · {_today()}*", "", "---", ""]
    for section in spec.get("sections") or []:
        if not isinstance(section, dict):
            continue
        if section.get("heading"):
            lines += [f"{'#' * min(int(section.get('level') or 2) + 1, 4)} {section['heading']}", ""]
        for block in section.get("blocks") or []:
            if isinstance(block, str):
                lines += [block, ""]
                continue
            kind = str(block.get("type") or "paragraph").lower()
            if kind in ("bullet", "bullets"):
                lines += [f"- {item}" for item in (block.get("items") or []) if str(item).strip()]
                lines.append("")
            elif kind in ("numbered", "steps"):
                lines += [f"{i}. {item}" for i, item in enumerate(block.get("items") or [], 1) if str(item).strip()]
                lines.append("")
            elif kind == "quote":
                lines += [f"> {block.get('text')}", ""]
            elif kind == "table":
                headers = block.get("headers") or []
                if headers:
                    lines.append("| " + " | ".join(str(h) for h in headers) + " |")
                    lines.append("|" + "|".join("---" for _ in headers) + "|")
                    for row in block.get("rows") or []:
                        lines.append("| " + " | ".join(_truncate(c, 80) for c in list(row)[: len(headers)]) + " |")
                    lines.append("")
            elif block.get("text"):
                lines += [str(block["text"]), ""]
    return "\n".join(lines).strip() + "\n"


HTML_TEMPLATE = """<!doctype html>
<html lang="{lang}"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title>
<style>
:root{{--gold:#c8a45c;--ink:#12141a;--paper:#fbfaf7;--muted:#6b7280}}
*{{box-sizing:border-box}}
body{{margin:0;background:linear-gradient(180deg,#0a0b10,#12141a);color:#f3f1ec;
font:16px/1.7 "Segoe UI",system-ui,-apple-system,sans-serif}}
.sheet{{max-width:900px;margin:48px auto;background:var(--paper);color:var(--ink);
border-radius:18px;overflow:hidden;box-shadow:0 40px 120px rgba(0,0,0,.55)}}
header{{padding:56px 56px 32px;border-top:4px solid var(--gold)}}
.kicker{{font-size:11px;letter-spacing:.22em;color:var(--gold);font-weight:700}}
h1{{font-size:38px;line-height:1.15;margin:14px 0 8px;font-weight:600}}
.sub{{color:var(--muted);font-size:16px;margin:0}}
main{{padding:12px 56px 56px}}
h2{{font-size:22px;margin:34px 0 10px;padding-bottom:8px;border-bottom:1px solid #e5dfd2}}
h3{{font-size:17px;margin:22px 0 8px;color:#3a3f4b}}
p{{margin:0 0 14px}}
ul,ol{{margin:0 0 16px;padding-left:22px}} li{{margin:5px 0}}
blockquote{{margin:18px 0;padding:14px 20px;border-left:3px solid var(--gold);
background:#f6f1e6;border-radius:0 10px 10px 0;font-style:italic}}
table{{width:100%;border-collapse:collapse;margin:18px 0;font-size:14px}}
th{{background:#1e222e;color:#e8c87e;text-align:left;padding:10px 12px;font-weight:600}}
td{{padding:9px 12px;border-bottom:1px solid #e8e3d8}}
tr:nth-child(even) td{{background:#faf8f3}}
footer{{padding:22px 56px;border-top:1px solid #e5dfd2;color:var(--muted);font-size:12px}}
@media print{{body{{background:#fff}}.sheet{{box-shadow:none;margin:0;border-radius:0}}}}
</style></head><body><article class="sheet">
<header><div class="kicker">{brand}</div><h1>{title}</h1>{subtitle}<p class="sub">{meta}</p></header>
<main>{body}</main>
<footer>{brand} · {date}</footer></article></body></html>
"""


def render_html(spec: dict[str, Any], *, brand: str = "Mestinghem AI", lang: str = "ru") -> str:
    import html as html_lib

    parts: list[str] = []
    if spec.get("summary"):
        parts.append(f"<blockquote>{html_lib.escape(str(spec['summary']))}</blockquote>")
    for section in spec.get("sections") or []:
        if not isinstance(section, dict):
            continue
        if section.get("heading"):
            tag = "h2" if int(section.get("level") or 1) <= 1 else "h3"
            parts.append(f"<{tag}>{html_lib.escape(str(section['heading']))}</{tag}>")
        for block in section.get("blocks") or []:
            if isinstance(block, str):
                parts.append(f"<p>{html_lib.escape(block)}</p>")
                continue
            kind = str(block.get("type") or "paragraph").lower()
            if kind in ("bullet", "bullets"):
                items = "".join(f"<li>{html_lib.escape(str(i))}</li>" for i in (block.get("items") or []) if str(i).strip())
                parts.append(f"<ul>{items}</ul>")
            elif kind in ("numbered", "steps"):
                items = "".join(f"<li>{html_lib.escape(str(i))}</li>" for i in (block.get("items") or []) if str(i).strip())
                parts.append(f"<ol>{items}</ol>")
            elif kind == "quote":
                parts.append(f"<blockquote>{html_lib.escape(str(block.get('text') or ''))}</blockquote>")
            elif kind == "table":
                headers = block.get("headers") or []
                if headers:
                    head = "".join(f"<th>{html_lib.escape(str(h))}</th>" for h in headers)
                    body = "".join(
                        "<tr>" + "".join(f"<td>{html_lib.escape(str(c))}</td>" for c in list(row)[: len(headers)]) + "</tr>"
                        for row in block.get("rows") or []
                    )
                    parts.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
            elif block.get("text"):
                parts.append(f"<p>{html_lib.escape(str(block['text']))}</p>")

    subtitle = f'<p class="sub">{html_lib.escape(str(spec["subtitle"]))}</p>' if spec.get("subtitle") else ""
    return HTML_TEMPLATE.format(
        lang=lang,
        brand=html_lib.escape(brand),
        title=html_lib.escape(str(spec.get("title") or "Документ")),
        subtitle=subtitle,
        meta=html_lib.escape(f"{spec.get('author') or brand}"),
        body="\n".join(parts),
        date=_today(),
    )

